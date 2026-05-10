import openpyxl as op
import time
import os

def type_out(text, speed):
    for char in text:
        print(char, end="", flush=True)
        time.sleep(speed)

def clear():
    os.system("cls")

clear()
order = "First", "Second", "Third"
counter = 0
headers = "ID", "First Name", "Last Name", "Birth Year", "Age"

book = op.Workbook()
sheet = book.active
sheet.append(headers)

while True:
    for item in order:
        type_out(f"{item} Favorite Person\n\n", .03)
        fname = input("First Name: ")
        lname = input("Last Name: ")
        year = input("Birth Year: ")
        print("")

        if not year.isdigit():
            clear()
            type_out("Birth Years must be number", .03)
            time.sleep(.3)
            type_out("\nInput not saved\n", .03)
            time.sleep(.3)
            type_out("Run the system again\n", .03)
            time.sleep(.3)
            type_out("Thank you!\n", .03)
            time.sleep(.3)
            exit()
            
        counter += 1
        birth_year = int(year)
        current_age = 2026 - birth_year
        
        sheet.append([f"0{counter}", fname, lname, birth_year, current_age])
        book.save("favorite_people.xlsx")
        clear()
    
    type_out("File saving...", .2)
    time.sleep(.2)
    clear()
    
    for char in "...":
        print(char, flush=True)
        time.sleep(.3)
    
    time.sleep(.2)
    clear()

    type_out("\n ---> Favorite People List <---\n\n", .03)
    
    load_book = op.load_workbook("favorite_people.xlsx")
    load_sheet = load_book.active
    
    for row in load_sheet.iter_rows(values_only=True):
        print(row)
        
    input_exit = input("\nPress Enter to Exit... ")
    clear()
    
    if input_exit == "":
        type_out("\nThank you for using my system", .03)
        break
    else:
        continue
